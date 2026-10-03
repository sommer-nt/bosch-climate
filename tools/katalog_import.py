"""Importiert die Produktdaten-Excel (Bosch Gesamtkatalog) in die Katalogdatei für Version 2.

Aufruf (aus dem Repo-Wurzelverzeichnis):
    python tools/katalog_import.py [pfad/zur/excel.xlsx]

Ohne Angabe wird splitklima/daten/quellen/Produktdaten_Split-Klima_Katalog_2026-03.xlsx gelesen.
Ergebnis: splitklima/daten/katalog.json – inklusive einer Liste von Prüfhinweisen zu
Datenlücken und Auffälligkeiten.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

WURZEL = Path(__file__).resolve().parent.parent
QUELLE = WURZEL / "splitklima" / "daten" / "quellen" / "Produktdaten_Split-Klima_Katalog_2026-03.xlsx"
ZIEL = WURZEL / "splitklima" / "daten" / "katalog.json"

FARBE = {"E": "weiß", "WE": "weiß", "F": "weiß", "ET": "anthrazit", "ES": "silber", "ER": "rot", "EB": "schwarz"}
BAUART = {"Wand": "Wandgerät", "Kassette": "Deckenkassette", "Konsole": "Konsole"}
# Heizleistung der CL3000iU-Inneneinheiten fehlt im Katalog; Werte aus Konfigurator v6.9.1
HEIZ_3000IU_V691 = {"CL3000iU W 20 E": 2.5, "CL3000iU W 26 E": 3.0, "CL3000iU W 35 E": 4.0, "CL3000iU W 53 E": 5.6}


def blatt(wb, name: str, kopfzeile: int = 4) -> list[dict]:
    ws = wb[name]
    zeilen = list(ws.iter_rows(values_only=True))
    kopf = [str(k).strip() if k is not None else "" for k in zeilen[kopfzeile - 1]]
    daten = []
    for z in zeilen[kopfzeile:]:
        if all(v is None for v in z):
            break  # Tabellenende (weitere Blöcke folgen ggf. nach Leerzeile)
        daten.append(dict(zip(kopf, z)))
    return daten


def zahl(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def leistung_aus_text(text) -> tuple[float | None, float | None]:
    """„2,5 / 3,2“ → (2.5, 3.2); „2,6 / –“ → (2.6, None); „–“ → (None, None)."""
    if text is None:
        return None, None
    teile = [t.strip() for t in str(text).split("/")]
    k = zahl(teile[0]) if teile else None
    h = zahl(teile[1]) if len(teile) > 1 else None
    return k, h


def klasse_aus_typ(typ: str) -> float | None:
    m = re.search(r"\b(\d{2})\b", typ.split("-Set")[-1] if "-Set" in typ else typ)
    return int(m.group(1)) / 10 if m else None


def farbe_aus_typ(typ: str) -> str:
    return FARBE.get(typ.split()[-1], "weiß")


def varianten(typ_tabelle: str) -> list[str]:
    """„CLC8001i-Set 25 E/ET/ES/ER“ → vier vollständige Typbezeichnungen.

    Nur ein Schrägstrich zwischen Buchstaben-Endungen trennt Varianten; „CL5000M 53/2 E“ bleibt ganz.
    """
    typ = typ_tabelle.strip()
    m = re.match(r"^(.*) ([A-Z]+(?:/[A-Z]+)+)$", typ)
    if not m:
        return [typ]
    return [f"{m.group(1)} {s}" for s in m.group(2).split("/")]


def verfuegbar(status: str) -> bool:
    return "ab Q3" not in (status or "")


def main(quelle: Path = QUELLE) -> dict:
    wb = openpyxl.load_workbook(quelle, data_only=True)
    hinweise: list[str] = []

    produkte = blatt(wb, "Katalog_Produkte")
    kennwerte = {}
    for k in blatt(wb, "Katalog_Kennwerte"):
        for typ in varianten(str(k["Typ (Katalogtabelle)"])):
            kennwerte[typ] = k
    sets_info = {s["Set"]: s for s in blatt(wb, "Kombi_Single_Sets")}
    matrix = blatt(wb, "Kombi_Matrix_Multi")
    ws = wb["Kombi_Matrix_Multi"]
    regeln = {}
    for z in ws.iter_rows(values_only=True):
        if z[0] and str(z[0]).startswith("CL") and isinstance(z[3], (int, float)) and "/" in str(z[0]):
            regeln[z[0]] = {"anschluesse": int(z[3]), "min_ie": int(z[4]), "einschraenkung": z[6]}

    def kw(typ, feld):
        k = kennwerte.get(typ)
        return zahl(k.get(feld)) if k else None

    def technik(typ: str, produkt: dict) -> dict:
        k_info, h_info = leistung_aus_text(produkt.get("Nennleistung K / H [kW] (Produktinfo)"))
        kuehl = kw(typ, "Kühlleistung nominal [kW]") or k_info
        heiz = kw(typ, "Heizleistung nominal [kW]") or h_info
        quelle = "Katalog"
        if kuehl is None:
            kuehl = klasse_aus_typ(typ)
            quelle = "aus Typbezeichnung abgeleitet"
            hinweise.append(f"{typ}: keine Kühlleistung im Katalog – {kuehl} kW aus der Typbezeichnung abgeleitet.")
        if heiz is None and typ in HEIZ_3000IU_V691:
            heiz = HEIZ_3000IU_V691[typ]
            hinweise.append(f"{typ}: keine Heizleistung im Katalog – {heiz} kW aus Konfigurator v6.9.1 übernommen.")
        if heiz is None:
            hinweise.append(f"{typ}: keine Heizleistung bekannt – wird für Heizbetrieb nicht empfohlen.")
        schall = kw(typ, "Schalldruck außen Kühlen [dB(A)]")
        if schall == 0:
            hinweise.append(f"{typ}: Schalldruck außen = 0 in der Excel – als „unbekannt“ behandelt.")
            schall = None
        return {
            "kuehl": kuehl, "kuehl_min": kw(typ, "Kühlleistung min. [kW]"), "kuehl_max": kw(typ, "Kühlleistung max. [kW]"),
            "heiz": heiz, "heiz_min": kw(typ, "Heizleistung min. [kW]"), "heiz_max": kw(typ, "Heizleistung max. [kW]"),
            "seer": kw(typ, "SEER / Arbeitszahl Kühlen"), "scop": kw(typ, "SCOP"),
            "schall_aussen": schall, "schall_innen_min": kw(typ, "Schalldruck innen niedrig [dB(A)]"),
            "schall_innen_max": kw(typ, "Schalldruck innen max. [dB(A)]"),
            "kaeltemittel": (kennwerte.get(typ) or {}).get("Kältemittel"),
            "datenquelle": quelle,
        }

    def basis(p: dict) -> dict:
        status = str(p.get("Lieferstatus") or "")
        return {
            "typ": p["Typ"], "bestellnr": str(p["Bestell-Nr."]), "linie": p["Produktlinie"],
            "beschreibung": p.get("Beschreibung"), "effizienz": p.get("Effizienzklasse K / H"),
            "preis": zahl(p["Preis netto [€] (UPE)"]), "lieferstatus": status, "verfuegbar": verfuegbar(status),
            "seite": p.get("Katalogseite"),
        }

    sets, aussen, innen = [], [], []
    matrix_nach_typ = {m["Inneneinheit"]: m for m in matrix}
    ae_spalten = [c for c in (matrix[0].keys() if matrix else []) if str(c).startswith("CL") and "/" in str(c)]
    for p in produkte:
        typ, rolle = p["Typ"], p["Rolle"]
        if rolle == "Set":
            info = sets_info.get(typ, {})
            hinweis = info.get("Hinweis")
            if hinweis and "Auslauf" in str(hinweis) and "Auslauf" not in str(p.get("Lieferstatus")):
                hinweise.append(f"{typ}: Lieferstatus „{p.get('Lieferstatus')}“, laut Set-Blatt aber „{hinweis}“.")
            sets.append({**basis(p), **technik(typ, p), "bauart": "Wandgerät", "farbe": farbe_aus_typ(typ),
                         "aussen_typ": info.get("Außeneinheit"), "innen_typ": info.get("Inneneinheit"),
                         "hinweis": hinweis})
        elif rolle == "Außeneinheit":
            r = regeln.get(typ, {})
            km = kennwerte.get(typ, {})
            km_anzahl = zahl(km.get("Max. Anzahl Inneneinheiten"))
            if r and km_anzahl and int(km_anzahl) != r["anschluesse"]:
                hinweise.append(f"{typ}: Kennwerte-Blatt nennt max. {int(km_anzahl)} Inneneinheiten, Regel-Blatt "
                                f"{r['anschluesse']} – Regel-Blatt verwendet.")
            aussen.append({**basis(p), **technik(typ, p), "anschluesse": r.get("anschluesse"),
                           "min_ie": r.get("min_ie", 1), "einschraenkung": r.get("einschraenkung")})
        else:
            m = matrix_nach_typ.get(typ, {})
            kompatibel = [c for c in ae_spalten if str(m.get(c, "")).strip() == "✔"]
            if not m:
                hinweise.append(f"{typ}: nicht in der Kombinationsmatrix – keine Außeneinheit zugeordnet.")
            if m.get("Hinweis") and "bestätigen" in str(m.get("Hinweis")):
                hinweise.append(f"{typ}: {m['Hinweis']}")
            bauart = next((v for k, v in BAUART.items() if k in rolle), rolle)
            innen.append({**basis(p), **technik(typ, p), "bauart": bauart, "farbe": farbe_aus_typ(typ),
                          "klasse": zahl(m.get("Leistungsklasse [kW]")) or klasse_aus_typ(typ),
                          "aussen_kompatibel": kompatibel})

    kombinationen = []
    for k in blatt(wb, "Katalog_Kombinationen"):
        klassen = [zahl(k[f"IE {i} [kW]"]) for i in range(1, 6) if zahl(k.get(f"IE {i} [kW]"))]
        kombinationen.append({"aussen": k["Außeneinheit"], "klassen": sorted(klassen, reverse=True),
                              "seite": k.get("Katalogseite")})

    zubehoer = [{"kategorie": z["Kategorie"], "name": z["Bezeichnung"], "bestellnr": str(z["Bestell-Nr."]),
                 "beschreibung": z.get("Beschreibung"), "preis": zahl(z["Preis netto [€]"]),
                 "passend": z.get("Passend für"), "seite": z.get("Katalogseite")}
                for z in blatt(wb, "Katalog_Zubehör_DL")]
    merkmale: dict[str, list[str]] = {}
    for z in blatt(wb, "Katalog_Merkmale"):
        merkmale.setdefault(z["Produktlinie"], []).append(z["Merkmal"])
    einsatz = [{"linie": z["Produktlinie"], "bereich": z["Bereich"], "merkmal": z["Merkmal"],
                "wert": z["Wert / Regel"], "seite": z.get("Katalogseite")} for z in blatt(wb, "Katalog_Einsatz_Montage")]

    katalog = {
        "quelle": "Bosch Gesamtkatalog Deutschland März 2026, Kapitel 3 Split-Klimageräte",
        "excel": quelle.name,
        "preise_freigegeben": True,
        "preisbasis": "Unverbindliche Preisempfehlung netto, zzgl. MwSt., Montage und Material",
        "sets": sets, "aussen": aussen, "innen": innen, "kombinationen": kombinationen,
        "zubehoer": zubehoer, "merkmale": merkmale, "einsatz": einsatz,
        "pruefhinweise": sorted(set(hinweise)),
    }
    ZIEL.write_text(json.dumps(katalog, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return katalog


if __name__ == "__main__":
    k = main(Path(sys.argv[1]) if len(sys.argv) > 1 else QUELLE)
    print(f"{len(k['sets'])} Sets, {len(k['aussen'])} Außeneinheiten, {len(k['innen'])} Inneneinheiten, "
          f"{len(k['kombinationen'])} Kombinationen, {len(k['zubehoer'])} Zubehör/DL → {ZIEL}")
    print(f"{len(k['pruefhinweise'])} Prüfhinweise:")
    for h in k["pruefhinweise"]:
        print(" -", h)
